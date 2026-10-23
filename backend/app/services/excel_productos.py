"""Importar y exportar productos en Excel.

Importar:
- Primero se PREVISUALIZA: cada fila dice si crea un producto, actualiza el precio de uno existente o tiene un error.
- Al confirmar se vuelve a analizar el archivo dentro de la transacción (los datos pudieron cambiar) y se guarda
  TODO junto: si algo falla a la mitad, no queda nada a medias. Las filas con error se omiten.
- Protecciones: tamaño máximo, número máximo de filas, tamaño descomprimido (un .xlsx es un zip) y lectura con
  defusedxml (openpyxl lo usa si está instalado).

Exportar: los textos que empiezan con = + - @ se escriben con un apóstrofo delante para que Excel no los ejecute
como fórmulas (inyección de fórmulas).
"""

import io
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.excepciones import ValidacionNegocio
from app.dependencies.auth import ContextoUsuario
from app.models import (
    Categoria,
    CodigoBarras,
    HistorialPrecio,
    Marca,
    MovimientoInventario,
    Producto,
    StockAlmacen,
)
from app.models.base import CERO
from app.models.enums import TipoAfectacionIGV, TipoPrecio, UnidadMedida
from app.schemas.inventario import StockInicialItem
from app.schemas.productos import ProductoCrear, normalizar_codigo
from app.services import auditoria, transacciones
from app.services.inventario import almacen_de_trabajo, aplicar_stock_inicial
from app.services.productos import almacen_venta_id, nuevo_producto, puede_ver_costos

TAMANIO_MAXIMO = 5 * 1024 * 1024
DESCOMPRIMIDO_MAXIMO = 60 * 1024 * 1024
FILAS_MAXIMAS = 5000
FILAS_EN_VISTA = 300

# columna normalizada -> (encabezado de la plantilla, ancho, ejemplo)
COLUMNAS: dict[str, tuple[str, int, Any]] = {
    "nombre": ("Nombre *", 34, "Gaseosa Inca Kola 1.5 L"),
    "codigo_barras": ("Código de barras", 18, "7750182006116"),
    "codigo_interno": ("Código interno", 14, ""),
    "categoria": ("Categoría", 18, "Bebidas"),
    "marca": ("Marca", 16, "Inca Kola"),
    "unidad": ("Unidad", 12, "UNIDAD"),
    "precio_venta": ("Precio venta *", 13, 6.50),
    "precio_costo": ("Precio costo", 13, 4.80),
    "stock_minimo": ("Stock mínimo", 12, 24),
    "igv": ("IGV", 12, "GRAVADO"),
    "maneja_lotes": ("Maneja lotes", 12, "NO"),
    "permite_decimales": ("Decimales", 11, "NO"),
    "stock_inicial": ("Stock inicial", 12, 48),
    "lote": ("Lote", 12, ""),
    "vencimiento": ("Vencimiento", 13, ""),
}
_ALIAS = {
    "nombre": "nombre",
    "producto": "nombre",
    "descripcion": "nombre",
    "codigo_de_barras": "codigo_barras",
    "codigo_barras": "codigo_barras",
    "codigo": "codigo_barras",
    "ean": "codigo_barras",
    "codigo_interno": "codigo_interno",
    "sku": "codigo_interno",
    "categoria": "categoria",
    "marca": "marca",
    "unidad": "unidad",
    "unidad_de_medida": "unidad",
    "precio_venta": "precio_venta",
    "precio": "precio_venta",
    "precio_de_venta": "precio_venta",
    "precio_costo": "precio_costo",
    "costo": "precio_costo",
    "precio_de_costo": "precio_costo",
    "stock_minimo": "stock_minimo",
    "minimo": "stock_minimo",
    "igv": "igv",
    "afectacion_igv": "igv",
    "maneja_lotes": "maneja_lotes",
    "lotes": "maneja_lotes",
    "decimales": "permite_decimales",
    "permite_decimales": "permite_decimales",
    "stock_inicial": "stock_inicial",
    "stock": "stock_inicial",
    "cantidad": "stock_inicial",
    "lote": "lote",
    "numero_lote": "lote",
    "vencimiento": "vencimiento",
    "fecha_vencimiento": "vencimiento",
}


def _normalizar(texto: Any) -> str:
    plano = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return "_".join(plano.strip().lower().replace("*", "").split())


def texto_seguro(valor: Any) -> Any:
    """Evita que un texto se interprete como fórmula al abrir el Excel."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


@dataclass
class Fila:
    numero: int
    datos: dict[str, Any]
    accion: str = "nuevo"  # nuevo | actualiza | sin_cambios | error
    errores: list[str] = field(default_factory=list)
    producto_id: int | None = None
    precio_anterior: Decimal | None = None
    crear: ProductoCrear | None = None
    stock: StockInicialItem | None = None


# --- Lectura -----------------------------------------------------------------------------------------------


def leer_archivo(contenido: bytes) -> list[tuple[int, dict[str, Any]]]:
    if len(contenido) > TAMANIO_MAXIMO:
        raise ValidacionNegocio("El archivo pesa más de 5 MB", codigo="archivo_muy_grande")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as comprimido:
            if sum(i.file_size for i in comprimido.infolist()) > DESCOMPRIMIDO_MAXIMO:
                raise ValidacionNegocio("El archivo es demasiado grande al descomprimirlo", codigo="archivo_muy_grande")
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except ValidacionNegocio:
        raise
    except Exception as exc:
        raise ValidacionNegocio(
            "No se pudo leer el archivo. Use la plantilla en formato Excel (.xlsx).", codigo="archivo_invalido"
        ) from exc
    try:
        hoja = libro.worksheets[0]
        filas = hoja.iter_rows(values_only=True)
        encabezado = next(filas, None)
        if not encabezado:
            raise ValidacionNegocio("El archivo está vacío", codigo="archivo_vacio")
        columnas = {i: _ALIAS.get(_normalizar(c)) for i, c in enumerate(encabezado)}
        presentes = {c for c in columnas.values() if c}
        faltan = {"nombre", "precio_venta"} - presentes
        if faltan:
            nombres = ", ".join(COLUMNAS[c][0].rstrip(" *") for c in sorted(faltan))
            raise ValidacionNegocio(f"Faltan columnas obligatorias: {nombres}", codigo="columnas_faltantes")
        resultado = []
        for numero, valores in enumerate(filas, start=2):
            if valores is None or all(v is None or str(v).strip() == "" for v in valores):
                continue
            if len(resultado) >= FILAS_MAXIMAS:
                raise ValidacionNegocio(
                    f"El archivo tiene más de {FILAS_MAXIMAS} productos: divídalo en varios", codigo="demasiadas_filas"
                )
            datos = {columnas[i]: v for i, v in enumerate(valores) if i in columnas and columnas[i]}
            resultado.append((numero, datos))
        if not resultado:
            raise ValidacionNegocio("El archivo no tiene productos", codigo="archivo_vacio")
        return resultado
    finally:
        libro.close()


def _texto(valor: Any, maximo: int) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # códigos numéricos que Excel guarda como número
    texto = " ".join(str(valor).split())
    if texto.startswith("'"):
        texto = texto[1:]
    return texto[:maximo] if texto else None


def _decimal(valor: Any, decimales: int) -> Decimal | None:
    if valor is None or str(valor).strip() == "":
        return None
    try:
        numero = Decimal(str(valor).replace(",", ".").replace("S/", "").strip())
    except InvalidOperation as exc:
        raise ValueError("no es un número") from exc
    if not numero.is_finite():
        raise ValueError("no es un número")
    return numero.quantize(Decimal(1).scaleb(-decimales))


def _si_no(valor: Any) -> bool:
    return _normalizar(valor) in {"si", "s", "x", "1", "true", "verdadero"}


def _fecha(valor: Any) -> date | None:
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(valor).strip(), formato).date()
        except ValueError:
            continue
    raise ValueError("fecha no válida (use DD/MM/AAAA)")


# --- Análisis ----------------------------------------------------------------------------------------------


async def analizar(session: AsyncSession, ctx: ContextoUsuario, contenido: bytes) -> list[Fila]:
    leidas = leer_archivo(contenido)
    filas = [Fila(numero=n, datos=d) for n, d in leidas]

    codigos_bd = dict((await session.execute(select(CodigoBarras.codigo, CodigoBarras.producto_id))).tuples().all())
    productos_bd = {p.id: p for p in await session.scalars(select(Producto).options(selectinload(Producto.codigos)))}
    por_nombre = {p.nombre.casefold(): p.id for p in productos_bd.values()}
    almacen_id = await almacen_venta_id(session, ctx.sucursal_id)
    con_movimientos = set(
        await session.scalars(
            select(MovimientoInventario.producto_id)
            .where(MovimientoInventario.almacen_id == (almacen_id or 0))
            .distinct()
        )
    )
    puede_stock = ctx.tiene("inventario.stock_inicial")

    vistos_codigo: dict[str, int] = {}
    vistos_nombre: dict[str, int] = {}
    for fila in filas:
        d = fila.datos
        errores = fila.errores
        nombre = _texto(d.get("nombre"), 200)
        if not nombre or len(nombre) < 2:
            errores.append("Falta el nombre")
        try:
            precio = _decimal(d.get("precio_venta"), 2)
            if precio is None:
                errores.append("Falta el precio de venta")
            elif precio < 0:
                errores.append("El precio no puede ser negativo")
            elif precio > Decimal("999999.99"):
                errores.append("El precio es demasiado alto")
        except ValueError:
            precio = None
            errores.append("El precio de venta no es un número")
        try:
            costo = _decimal(d.get("precio_costo"), 4) or CERO
            if costo < 0:
                errores.append("El costo no puede ser negativo")
        except ValueError:
            costo = CERO
            errores.append("El costo no es un número")
        codigo = None
        if (crudo := _texto(d.get("codigo_barras"), 60)) is not None:
            try:
                codigo = normalizar_codigo(crudo)
            except ValueError:
                errores.append("Código de barras no válido")
        unidad = (_texto(d.get("unidad"), 20) or "UNIDAD").upper()
        if unidad not in UnidadMedida.__members__:
            errores.append(f"Unidad «{unidad}» no válida")
            unidad = "UNIDAD"
        igv = (_texto(d.get("igv"), 20) or "GRAVADO").upper()
        if igv not in TipoAfectacionIGV.__members__:
            errores.append("IGV debe ser GRAVADO, EXONERADO o INAFECTO")
            igv = "GRAVADO"
        maneja_lotes = _si_no(d.get("maneja_lotes"))
        decimales = _si_no(d.get("permite_decimales")) or unidad in ("KILOGRAMO", "LITRO")
        try:
            minimo = _decimal(d.get("stock_minimo"), 3) or CERO
            cantidad = _decimal(d.get("stock_inicial"), 3)
            vencimiento = _fecha(d.get("vencimiento"))
        except ValueError as exc:
            errores.append(f"Stock o vencimiento: {exc}")
            minimo, cantidad, vencimiento = CERO, None, None
        if minimo < 0 or (cantidad is not None and cantidad < 0):
            errores.append("Las cantidades no pueden ser negativas")
        for valor in (minimo, cantidad):
            if valor is not None and not decimales and valor != valor.to_integral_value():
                errores.append("Cantidad con decimales en un producto que no los admite")
                break
        lote = _texto(d.get("lote"), 50)

        # ¿Existe? Por código de barras y si no, por nombre exacto
        existente_id = codigos_bd.get(codigo) if codigo else None
        if existente_id is None and nombre:
            existente_id = por_nombre.get(nombre.casefold())

        if codigo:
            if codigo in vistos_codigo:
                errores.append(f"Código repetido: es el de la fila {vistos_codigo[codigo]}")
            else:
                vistos_codigo[codigo] = fila.numero
        if nombre:
            clave = nombre.casefold()
            if clave in vistos_nombre:
                errores.append(f"Producto repetido: es el de la fila {vistos_nombre[clave]}")
            else:
                vistos_nombre[clave] = fila.numero

        producto = productos_bd.get(existente_id) if existente_id else None
        if cantidad:
            lote_producto = producto.maneja_lotes if producto else maneja_lotes
            if not puede_stock:
                errores.append("No tiene permiso para cargar stock inicial")
            elif producto and producto.id in con_movimientos:
                errores.append("Ya tiene movimientos: su stock se cambia con un ajuste")
            elif lote_producto and not lote:
                errores.append("Indique el lote del stock inicial")
            elif not lote_producto and lote:
                errores.append("El producto no maneja lotes")

        if errores:
            fila.accion = "error"
            continue
        if producto:
            fila.producto_id = producto.id
            fila.precio_anterior = producto.precio_venta
            fila.accion = "actualiza" if producto.precio_venta != precio or cantidad else "sin_cambios"
            if cantidad and not producto.permite_decimales and cantidad != cantidad.to_integral_value():
                fila.accion, fila.errores = "error", ["Cantidad con decimales en un producto que no los admite"]
                continue
        else:
            fila.crear = ProductoCrear(
                nombre=nombre,
                codigos=[codigo] if codigo else [],
                codigo_interno=_texto(d.get("codigo_interno"), 30),
                unidad_medida=UnidadMedida(unidad),
                precio_venta=precio,
                precio_costo=costo if not cantidad else CERO,  # con stock, el costo lo fija la carga
                stock_minimo=minimo,
                tipo_afectacion_igv=TipoAfectacionIGV(igv),
                maneja_lotes=maneja_lotes,
                permite_decimales=decimales,
            )
            fila.datos["_categoria"] = _texto(d.get("categoria"), 80)
            fila.datos["_marca"] = _texto(d.get("marca"), 80)
        if cantidad:
            fila.stock = StockInicialItem(
                producto_id=producto.id if producto else 0,
                cantidad=cantidad,
                costo_unitario=costo,
                numero_lote=lote,
                fecha_vencimiento=vencimiento,
            )
        fila.datos["_precio"] = precio
    return filas


def resumen_vista(filas: list[Fila]) -> dict[str, Any]:
    conteo = {"nuevo": 0, "actualiza": 0, "sin_cambios": 0, "error": 0}
    for f in filas:
        conteo[f.accion] += 1
    # Primero los errores (lo que hay que corregir) y luego el resto, con un límite para la pantalla
    ordenadas = sorted(filas, key=lambda f: (f.accion != "error", f.numero))[:FILAS_EN_VISTA]
    return {
        "total": len(filas),
        "nuevos": conteo["nuevo"],
        "actualizan": conteo["actualiza"],
        "sin_cambios": conteo["sin_cambios"],
        "errores": conteo["error"],
        "con_stock": sum(1 for f in filas if f.stock and f.accion != "error"),
        "filas": [
            {
                "fila": f.numero,
                "nombre": _texto(f.datos.get("nombre"), 200),
                "codigo_barras": _texto(f.datos.get("codigo_barras"), 60),
                "precio_venta": f.datos.get("_precio"),
                "precio_anterior": f.precio_anterior,
                "stock_inicial": f.stock.cantidad if f.stock else None,
                "accion": f.accion,
                "errores": f.errores,
            }
            for f in ordenadas
        ],
    }


# --- Importación -------------------------------------------------------------------------------------------


async def _por_nombre(session: AsyncSession, modelo, nombre: str | None, cache: dict[str, int]) -> int | None:
    if not nombre:
        return None
    clave = nombre.casefold()
    if clave not in cache:
        objeto = await session.scalar(select(modelo).where(modelo.nombre == nombre))
        if objeto is None:
            objeto = modelo(nombre=nombre)
            session.add(objeto)
            await session.flush()
        cache[clave] = objeto.id
    return cache[clave]


async def importar(session: AsyncSession, ctx: ContextoUsuario, contenido: bytes, clave: str | None) -> dict[str, Any]:
    resultado: dict[str, Any] = {}

    async def accion() -> int:
        filas = await analizar(session, ctx, contenido)
        validas = [f for f in filas if f.accion in ("nuevo", "actualiza")]
        if not validas:
            raise ValidacionNegocio("No hay filas válidas para importar", codigo="sin_filas_validas")
        categorias: dict[str, int] = {}
        marcas: dict[str, int] = {}
        stock_items: list[StockInicialItem] = []
        creados = actualizados = 0
        for f in validas:
            if f.crear is not None:
                producto = nuevo_producto(f.crear)
                producto.categoria_id = await _por_nombre(session, Categoria, f.datos.get("_categoria"), categorias)
                producto.marca_id = await _por_nombre(session, Marca, f.datos.get("_marca"), marcas)
                session.add(producto)
                await session.flush()
                creados += 1
                producto_id = producto.id
            else:
                producto = await session.scalar(
                    select(Producto).where(Producto.id == f.producto_id).with_for_update(of=Producto)
                )
                nuevo_precio = f.datos["_precio"]
                if producto.precio_venta != nuevo_precio:
                    session.add(
                        HistorialPrecio(
                            producto_id=producto.id,
                            tipo=TipoPrecio.VENTA,
                            precio_anterior=producto.precio_venta,
                            precio_nuevo=nuevo_precio,
                            origen="IMPORTACION",
                            usuario_id=ctx.usuario_id,
                        )
                    )
                    producto.precio_venta = nuevo_precio
                    producto.version += 1
                actualizados += 1
                producto_id = producto.id
            if f.stock is not None:
                stock_items.append(f.stock.model_copy(update={"producto_id": producto_id}))

        stock_resultado = None
        if stock_items:
            almacen = await almacen_de_trabajo(session, ctx, None)
            stock_resultado = await aplicar_stock_inicial(session, ctx, almacen, stock_items, origen="IMPORTACION")
        resultado.update(
            creados=creados,
            actualizados=actualizados,
            omitidos=sum(1 for f in filas if f.accion == "error"),
            con_stock=len(stock_items),
            valor_stock=stock_resultado["valor"] if stock_resultado else CERO,
        )
        auditoria.registrar(
            session,
            accion="IMPORTAR",
            modulo="productos",
            usuario_id=ctx.usuario_id,
            sucursal_id=ctx.sucursal_id,
            descripcion=f"Importó productos desde Excel: {creados} nuevos, {actualizados} actualizados",
            datos={k: str(v) for k, v in resultado.items()},
            ip=ctx.ip,
        )
        return creados + actualizados

    await transacciones.ejecutar(
        session, accion, usuario_id=ctx.usuario_id, operacion="importar_productos", clave=clave
    )
    return resultado or {
        "creados": 0,
        "actualizados": 0,
        "omitidos": 0,
        "con_stock": 0,
        "valor_stock": CERO,
        "repetida": True,
    }


# --- Plantilla y exportación -------------------------------------------------------------------------------


def _libro(encabezados: list[tuple[str, int]]) -> tuple[Workbook, Any]:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Productos"
    negrita = Font(bold=True, color="14213D")
    fondo = PatternFill("solid", fgColor="FFC300")
    for i, (titulo, ancho) in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=i, value=titulo)
        celda.font, celda.fill = negrita, fondo
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    return libro, hoja


def _bytes(libro: Workbook) -> bytes:
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def plantilla() -> bytes:
    libro, hoja = _libro([(t, a) for t, a, _ in COLUMNAS.values()])
    hoja.append([texto_seguro(e) for _, _, e in COLUMNAS.values()])
    ayuda = libro.create_sheet("Instrucciones")
    for linea in (
        "Columnas obligatorias: Nombre y Precio venta. Las demás son opcionales.",
        "Si el código de barras (o el nombre) ya existe, solo se actualiza el precio de venta.",
        "Unidad: UNIDAD, KILOGRAMO, GRAMO, LITRO, MILILITRO, PAQUETE, CAJA o DOCENA.",
        "IGV: GRAVADO, EXONERADO (frutas, verduras, huevos...) o INAFECTO.",
        "Maneja lotes / Decimales: SI o NO. Kilos y litros admiten decimales automáticamente.",
        "Stock inicial: cantidad que hay hoy en el almacén de venta.",
        "Si el producto maneja lotes, indique también Lote y Vencimiento (DD/MM/AAAA).",
        "Las categorías y marcas que no existan se crean solas.",
        f"Máximo {FILAS_MAXIMAS} productos por archivo y 5 MB.",
    ):
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 110
    return _bytes(libro)


async def exportar(session: AsyncSession, ctx: ContextoUsuario, *, activo: bool | None = True) -> bytes:
    almacen_id = await almacen_venta_id(session, ctx.sucursal_id)
    ver_costos = puede_ver_costos(ctx)
    columnas = ["nombre", "codigo_barras", "codigo_interno", "categoria", "marca", "unidad", "precio_venta"]
    if ver_costos:
        columnas.append("precio_costo")
    columnas += ["stock_minimo", "igv", "maneja_lotes", "permite_decimales"]
    encabezados = [(COLUMNAS[c][0].rstrip(" *"), COLUMNAS[c][1]) for c in columnas] + [("Stock actual", 12)]
    libro, hoja = _libro(encabezados)
    consulta = (
        select(Producto, func.coalesce(StockAlmacen.cantidad, 0))
        .outerjoin(
            StockAlmacen,
            (StockAlmacen.producto_id == Producto.id) & (StockAlmacen.almacen_id == (almacen_id or 0)),
        )
        .options(selectinload(Producto.codigos))
        .order_by(Producto.nombre)
    )
    if activo is not None:
        consulta = consulta.where(Producto.activo.is_(activo))
    for producto, cantidad in (await session.execute(consulta)).tuples():
        base = [c.codigo for c in producto.codigos if c.presentacion_id is None]
        valores = {
            "nombre": producto.nombre,
            "codigo_barras": base[0] if base else None,
            "codigo_interno": producto.codigo_interno,
            "categoria": producto.categoria.nombre if producto.categoria else None,
            "marca": producto.marca.nombre if producto.marca else None,
            "unidad": producto.unidad_medida.value,
            "precio_venta": float(producto.precio_venta),
            "precio_costo": float(producto.precio_costo),
            "stock_minimo": float(producto.stock_minimo),
            "igv": producto.tipo_afectacion_igv.value,
            "maneja_lotes": "SI" if producto.maneja_lotes else "NO",
            "permite_decimales": "SI" if producto.permite_decimales else "NO",
        }
        hoja.append([texto_seguro(valores[c]) for c in columnas] + [float(cantidad)])
    return _bytes(libro)
